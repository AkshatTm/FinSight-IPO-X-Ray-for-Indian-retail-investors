"""An Excel sheet for labelling gold v1, and the converter back to ``gold_values.jsonl``.

Labelling 110 values in raw JSON is error-prone, so Akshat fills a sheet instead:

    uv run python -m finsight.evaluate.gold export-xlsx   # writes data/gold/gold_labelling.xlsx
    uv run python -m finsight.evaluate.gold import-xlsx   # sheet -> gold_values.jsonl + checks

The sheet is a working file (git-ignored); only the validated ``gold_values.jsonl`` is committed.
Mistakes are reported in plain English: row number, what is wrong, how to fix it.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from finsight.core.config import get_settings
from finsight.evaluate.gold import (
    DEFAULT_DOC,
    FIELDS,
    template_rows,
    validate_file,
    validate_row,
)
from finsight.ingest.registry import list_demo_ipos

HEADERS = ["ipo_id", "field_id", "doc", "value_raw", "page", "quote", "status", "notes"]
_COL = {name: i + 1 for i, name in enumerate(HEADERS)}
LIST_SEP, PAIR_SEP = " | ", " :: "

WHAT_TO_FIND = {
    "fresh_issue_size": "Rupee size of the Fresh Issue (new shares). Cover page or 'The Offer'.",
    "ofs_shares": "Number of shares in the Offer for Sale (existing holders selling).",
    "ofs_amount": "Rupee size of the Offer for Sale. Often [●] in the RHP.",
    "offer_price": "Final price per share. Only the final Prospectus states it (RHP: [●]).",
    "price_band": "Low and high price, e.g. ₹ 103 to ₹ 108. Not in document if a fixed price.",
    "total_issue_size": "Total size of the offer in rupees. Final Prospectus (RHP: [●]).",
    "face_value": "Face value per share, e.g. ₹ 1 or ₹ 10.",
    "book_running_lead_managers": "All BRLMs (the lead banks), as named on the cover.",
    "registrar": "Registrar to the Offer.",
    "promoters": "Names of the promoters ('Our Promoters' / cover).",
    "objects_of_offer": "Each use of the Fresh Issue money and its amount ('Objects of the Offer').",
}

# One worked example per value type (validated by tests; numbers are illustrative only).
EXAMPLES: list[dict[str, Any]] = [
    {
        "type": "money",
        "field_id": "fresh_issue_size",
        "value_raw": "₹ 4,720 million",
        "page": 163,
        "quote": "a Fresh Issue of [●] Equity Shares aggregating up to ₹ 4,720 million by our Company",
        "status": "present",
        "note": "Copy the amount and unit exactly as printed.",
    },
    {
        "type": "count",
        "field_id": "ofs_shares",
        "value_raw": "10,181,585",
        "page": 3,
        "quote": "Offer for Sale of up to 10,181,585 Equity Shares of face value of ₹ 1 each",
        "status": "present",
        "note": "Just the number of shares.",
    },
    {
        "type": "range",
        "field_id": "price_band",
        "value_raw": "₹ 103 to ₹ 108",
        "page": 2,
        "quote": "Price Band: ₹ 103 to ₹ 108 per Equity Share",
        "status": "present",
        "note": "Both ends.",
    },
    {
        "type": "text",
        "field_id": "registrar",
        "value_raw": "Link Intime India Private Limited",
        "page": 8,
        "quote": "Registrar to the Offer Link Intime India Private Limited",
        "status": "present",
        "note": "One name.",
    },
    {
        "type": "list",
        "field_id": "book_running_lead_managers",
        "value_raw": "Kotak Mahindra Capital Company Limited | Morgan Stanley India Company Private Limited",
        "page": 1,
        "quote": "Kotak Mahindra Capital Company Limited Morgan Stanley India Company Private Limited",
        "status": "present",
        "note": 'Separate names with " | " (space, bar, space).',
    },
    {
        "type": "table",
        "field_id": "objects_of_offer",
        "value_raw": "Expenditure for new technology :: 1,900.00 | General corporate purposes :: [●]",
        "page": 164,
        "quote": "Expenditure for new technology 1,900.00 General corporate purposes [●]",
        "status": "present",
        "note": 'Each row is "purpose :: amount"; separate rows with " | ". Amount as printed.',
    },
    {
        "type": "placeholder",
        "field_id": "ofs_amount",
        "value_raw": "[●]",
        "page": 3,
        "quote": "an Offer for Sale aggregating up to ₹ [●] million",
        "status": "placeholder",
        "note": "The document leaves a blank: type [●] and pick 'placeholder'.",
    },
    {
        "type": "not in document",
        "field_id": "price_band",
        "value_raw": "",
        "page": "",
        "quote": "",
        "status": "not_in_document",
        "note": "The document does not state it: leave value, page and quote empty.",
    },
]


def export_xlsx(path: Path, force: bool = False) -> Path:
    """Write the labelling sheet (only ipo_id, field_id and doc are pre-filled)."""
    if path.exists() and not force:
        raise FileExistsError(f"{path} already exists (your labels may be in it); use --force")
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "LABELS"
    ws.append(HEADERS)
    for row in template_rows():
        ws.append([row["ipo_id"], row["field_id"], row["doc"]])
    head, grey = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="E7E6E6")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
    for r in range(2, ws.max_row + 1):
        for c in range(1, 4):
            ws.cell(r, c).fill = grey
        for name in ("value_raw", "quote", "notes"):
            ws.cell(r, _COL[name]).alignment = Alignment(wrap_text=True, vertical="top")
    for name, width in zip(HEADERS, (26, 28, 11, 38, 7, 60, 16, 30), strict=True):
        ws.column_dimensions[get_column_letter(_COL[name])].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"
    dropdown = DataValidation(
        type="list", formula1='"present,placeholder,not_in_document"', allow_blank=True
    )
    dropdown.error, dropdown.errorTitle = "Pick one from the list", "Status"
    dropdown.add(f"G2:G{ws.max_row}")
    ws.add_data_validation(dropdown)
    _how_to(wb.create_sheet("HOW TO"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _how_to(ws: Any) -> None:
    bold = Font(bold=True)
    steps = [
        "HOW TO FILL THE LABELS SHEET",
        "1. Work on the LABELS sheet. One row = one value to find. Do not change columns A-C (grey).",
        "2. Open the PDF named in the table at the bottom of this sheet (rhp or prospectus, per column C).",
        "3. Label BLIND: read the value from the PDF yourself; do not run any extractor first.",
        "4. value_raw = the value exactly as printed (keep the ₹, the unit word, the commas).",
        "5. page = the PDF page, i.e. the number your PDF viewer shows, NOT the small number printed on the page.",
        "6. quote = copy-paste the sentence or table line where you found it. The value must appear inside it.",
        "7. status = pick from the dropdown (see below). notes = anything unusual (optional).",
        "8. Save, then run:  uv run python -m finsight.evaluate.gold import-xlsx   (add --partial while you are halfway)",
        "",
        "STATUS",
        "present          the document states the value",
        "placeholder      the document leaves a blank, printed as [●]  (type [●] as the value)",
        "not_in_document  the document does not state it at all (leave value, page, quote empty)",
        "",
        "EXAMPLES, one per value type (these numbers are made up)",
    ]
    for line in steps:
        ws.append([line])
        if line.isupper() or line.startswith("EXAMPLES") or line == "HOW TO FILL THE LABELS SHEET":
            ws.cell(ws.max_row, 1).font = bold
    ws.append(["type", "field_id", "value_raw", "page", "quote", "status", "tip"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for ex in EXAMPLES:
        ws.append([ex["type"], ex["field_id"], ex["value_raw"], ex["page"], ex["quote"],
                   ex["status"], ex["note"]])  # fmt: skip
    ws.append([])
    ws.append(["THE 11 FIELDS"])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["field_id", "type", "read it from", "what to look for"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for field_id, kind in FIELDS.items():
        ws.append([field_id, kind, DEFAULT_DOC[field_id], WHAT_TO_FIND[field_id]])
    ws.append([])
    ws.append(["PDF FILES"])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["ipo_id", "rhp file", "prospectus file"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for ipo in list_demo_ipos():
        ws.append([ipo.ipo_id, ipo.rhp.file.name, ipo.prospectus.file.name])
    ws.column_dimensions["A"].width = 28
    for col, width in zip("BCDEFG", (30, 44, 8, 60, 16, 50), strict=True):
        ws.column_dimensions[col].width = width


# ------------------------------------------------------------------------------ import
_TIPS = {
    "money": "Write it as printed, with the currency and unit, e.g. ₹ 4,720 million or ₹ 800.00 crore.",
    "count": "Write the number of shares as printed, e.g. 10,181,585.",
    "range": "Write both ends, e.g. ₹ 103 to ₹ 108.",
}


def friendly(problem: str, field_id: str, value: Any) -> str:
    """Plain-English version of a validator message, including how to fix it."""
    kind = FIELDS.get(field_id, "text")
    if problem.startswith(("unknown ipo_id", "unknown field_id", "doc must be")):
        return "Columns A-C (ipo_id, field_id, doc) were changed. Restore the original text."
    if problem.startswith("status must be"):
        return ("Status is missing or not one of the allowed words. Pick present, placeholder or "
                "not_in_document from the dropdown.")  # fmt: skip
    if problem.startswith("not_in_document rows"):
        return ("Status is not_in_document, so value, page and quote must be empty. Clear them, "
                "or change the status to present.")  # fmt: skip
    if problem.startswith("value_raw is empty"):
        return ("Value is empty. Copy the value exactly as printed in the PDF, or set the status "
                "to not_in_document if the document does not state it.")  # fmt: skip
    if problem.startswith("page must be"):
        return ("Page must be a whole number: the PDF page shown by your PDF viewer (not the small "
                "number printed on the page).")  # fmt: skip
    if m := re.match(r"page (\d+) is beyond the document \((\d+) pages\)", problem):
        return (f"Page {m[1]} is past the end of this document ({m[2]} pages). Use the page number "
                "shown by your PDF viewer.")  # fmt: skip
    if problem.startswith("quote is empty"):
        return "Quote is empty. Paste the sentence or table line from the PDF where you found it."
    if problem.startswith("value_raw does not appear"):
        return (
            "The value is not inside the quote, so it cannot be checked. Copy the value from "
            "the quote text (spaces and line breaks do not matter) or paste a longer quote."
        )
    if problem.startswith("value_raw") and "does not parse as" in problem:
        noun = {"money": "an amount", "count": "a number of shares", "range": "a price band"}
        return f"Cannot read {value!r} as {noun.get(kind, kind)}. " + _TIPS.get(kind, "")
    if problem.startswith("placeholder:"):
        return "Status is placeholder, so the value must be the blank as printed: [●]."
    if problem.startswith("value is [●]"):
        return "The value is a blank ([●]). Set the status to placeholder."
    if "list of text" in problem:
        return 'Write the names separated by " | " (space, bar, space), e.g. Kotak Mahindra | Morgan Stanley.'
    if "pairs" in problem:
        return ('Write each row as "purpose :: amount" and separate rows with " | ", e.g. '
                "Capex :: ₹ 1,900.00 million | General corporate purposes :: [●].")  # fmt: skip
    return problem


def _text(cell: Any) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def _page(cell: Any) -> Any:
    if cell in (None, ""):
        return None
    if isinstance(cell, int | float) and float(cell).is_integer():
        return int(cell)
    return int(cell) if str(cell).strip().isdigit() else str(cell)


def _value(kind: str, text: str) -> Any:
    if kind == "list":
        return [p.strip() for p in text.split("|") if p.strip()]
    if kind == "table":
        rows = []
        for item in (p.strip() for p in text.split("|") if p.strip()):
            parts = [x.strip() for x in item.split("::")]
            rows.append(parts if len(parts) == 2 and all(parts) else [item])
        return rows
    return text


@dataclass
class ImportResult:
    written: int = 0
    unfilled: int = 0
    errors: list[str] = field(default_factory=list)


def import_xlsx(xlsx: Path, out: Path, partial: bool = False) -> ImportResult:
    """Sheet -> ``out`` (JSONL). Writes nothing unless every checked row is valid."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["LABELS"] if "LABELS" in wb.sheetnames else wb.worksheets[0]
    result = ImportResult()
    if [_text(c.value) for c in ws[1]][: len(HEADERS)] != HEADERS:
        result.errors.append(
            "Row 1: the column headings were changed. They must be: " + ", ".join(HEADERS) + "."
        )
        return result
    today = date.today().isoformat()
    rows: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        cells = {name: ws.cell(r, col).value for name, col in _COL.items()}
        ipo_id, field_id, doc = (_text(cells[k]) for k in ("ipo_id", "field_id", "doc"))
        if not (ipo_id or field_id):
            continue
        where = f"Row {r} ({ipo_id} — {field_id})"
        if not any(_text(cells[k]) for k in ("value_raw", "page", "quote", "status")):
            result.unfilled += 1
            if not partial:
                result.errors.append(
                    f"{where}: not filled in yet. Fill value, page, quote and status, or set the "
                    "status to not_in_document if the document does not state it."
                )
            continue
        status = _text(cells["status"]).lower()
        kind = FIELDS.get(field_id, "text")
        raw_value = _text(cells["value_raw"])
        row = {
            "ipo_id": ipo_id, "field_id": field_id, "doc": doc,
            "value_raw": "" if status == "not_in_document" and not raw_value
            else _value(kind, raw_value),
            "page": _page(cells["page"]), "quote": _text(cells["quote"]), "status": status,
            "labelled_at": today, "notes": _text(cells["notes"]),
        }  # fmt: skip
        if problems := validate_row(row):
            fixes = " Also: ".join(friendly(p, field_id, raw_value) for p in problems)
            result.errors.append(f"{where}: {fixes}")
        rows.append(row)
    if result.errors:
        return result
    out.parent.mkdir(parents=True, exist_ok=True)
    lines = [json.dumps(row, ensure_ascii=False) for row in rows]
    out.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8", newline="\n")
    leftover = validate_file(out, require_complete=not partial)
    if leftover:  # should not happen after row checks; never leave a bad file behind
        out.unlink()
        result.errors += [f"Final check failed: {issue}" for issue in leftover]
        return result
    result.written = len(rows)
    return result


def default_paths() -> tuple[Path, Path]:
    gold_dir = get_settings().paths.gold_dir
    return gold_dir / "gold_labelling.xlsx", gold_dir / "gold_values.jsonl"
