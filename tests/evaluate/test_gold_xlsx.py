import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from finsight.evaluate.gold import FIELDS, validate_file
from finsight.evaluate.gold_xlsx import HEADERS, export_xlsx, import_xlsx

URBAN = "urban-company-2025"


def _sheet(path: Path) -> Any:
    return openpyxl.load_workbook(path)["LABELS"]


def _find(ws: Any, ipo: str, field: str) -> int:
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == ipo and ws.cell(r, 2).value == field:
            return int(r)
    raise AssertionError((ipo, field))


def _fill(ws: Any, ipo: str, field: str, **cells: Any) -> int:
    r = _find(ws, ipo, field)
    for name, value in cells.items():
        ws.cell(r, HEADERS.index(name) + 1, value)
    return r


@pytest.fixture
def xlsx(tmp_path: Path) -> Path:
    return export_xlsx(tmp_path / "gold_labelling.xlsx")


def test_export_layout(xlsx: Path) -> None:
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["LABELS", "HOW TO"]
    ws = wb["LABELS"]
    assert [c.value for c in ws[1]] == HEADERS == [
        "ipo_id", "field_id", "doc", "value_raw", "page", "quote", "status", "notes",
    ]  # fmt: skip
    assert ws.max_row == 1 + 110
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[0]
        assert row[1]
        assert row[2] in ("rhp", "prospectus")
        assert all(v is None for v in row[3:]), "only ipo_id, field_id and doc are pre-filled"
    assert ws.freeze_panes == "A2"


def test_status_column_has_a_dropdown(xlsx: Path) -> None:
    ws = _sheet(xlsx)
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    dv = validations[0]
    assert dv.type == "list"
    assert dv.formula1 == '"present,placeholder,not_in_document"'
    assert "G2" in str(dv.sqref)


def test_how_to_sheet_has_one_example_per_type(xlsx: Path) -> None:
    text = "\n".join(
        str(c.value) for row in openpyxl.load_workbook(xlsx)["HOW TO"].iter_rows() for c in row
        if c.value is not None
    )  # fmt: skip
    for kind in ("money", "count", "range", "text", "list", "table"):
        assert kind in text
    assert " | " in text
    assert " :: " in text
    for field in FIELDS:
        assert field in text
    assert "not_in_document" in text
    assert "PDF page" in text


def test_export_never_overwrites_a_filled_sheet(xlsx: Path) -> None:
    with pytest.raises(FileExistsError, match="--force"):
        export_xlsx(xlsx)
    assert export_xlsx(xlsx, force=True) == xlsx


def test_import_converts_every_field_type(xlsx: Path, tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["LABELS"]
    _fill(ws, URBAN, "fresh_issue_size", value_raw="₹ 4,720 million", page=163,
          quote="Fresh Issue ... aggregating up to ₹ 4,720 million by our Company",
          status="present")  # fmt: skip
    _fill(
        ws,
        URBAN,
        "ofs_shares",
        value_raw=12345678,
        page=3.0,
        quote="12345678 Equity Shares",
        status="present",
    )  # numbers typed into cells arrive as int/float
    _fill(ws, URBAN, "ofs_amount", value_raw="[●]", page=3, status="placeholder",
          quote="an Offer for Sale aggregating up to ₹ [●] million")  # fmt: skip
    _fill(
        ws,
        URBAN,
        "registrar",
        value_raw="Link Intime India",
        page=5,
        quote="Registrar to the Offer Link Intime India Private Limited",
        status="present",
    )
    _fill(
        ws,
        URBAN,
        "book_running_lead_managers",
        value_raw="Kotak | Morgan Stanley | ",
        page=5,
        quote="Kotak Mahindra Capital Morgan Stanley India Company",
        status="present",
    )
    _fill(ws, URBAN, "objects_of_offer",
          value_raw="Capex :: ₹ 1,900.00 million | General corporate purposes :: [●]", page=164,
          quote="Capex 1,900.00 General corporate purposes [●] ₹ 1,900.00 million",
          status="present")  # fmt: skip
    _fill(ws, URBAN, "price_band", status="not_in_document")
    wb.save(xlsx)
    out = tmp_path / "gold_values.jsonl"
    result = import_xlsx(xlsx, out, partial=True)
    assert result.errors == []
    assert result.written == 7
    rows = {
        json.loads(line)["field_id"]: json.loads(line)
        for line in out.read_text("utf-8").splitlines()
    }
    assert rows["ofs_shares"]["value_raw"] == "12345678"
    assert rows["ofs_shares"]["page"] == 3
    assert rows["book_running_lead_managers"]["value_raw"] == ["Kotak", "Morgan Stanley"]
    assert rows["objects_of_offer"]["value_raw"] == [
        ["Capex", "₹ 1,900.00 million"], ["General corporate purposes", "[●]"],
    ]  # fmt: skip
    assert rows["price_band"]["status"] == "not_in_document"
    assert rows["price_band"]["value_raw"] == ""
    assert rows["price_band"]["page"] is None
    assert rows["fresh_issue_size"]["labelled_at"]
    assert validate_file(out) == []


def test_default_import_needs_every_row_and_writes_nothing_if_invalid(
    xlsx: Path, tmp_path: Path
) -> None:
    out = tmp_path / "gold_values.jsonl"
    result = import_xlsx(xlsx, out)
    assert result.written == 0
    assert not out.exists()
    assert len(result.errors) == 110
    assert result.errors[0].startswith("Row 2 (")
    assert "not filled in yet" in result.errors[0]


def test_errors_are_plain_english_with_row_and_fix(xlsx: Path, tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["LABELS"]
    r1 = _fill(ws, URBAN, "fresh_issue_size", value_raw="four thousand", page=163,
               quote="four thousand", status="present")  # fmt: skip
    r2 = _fill(ws, URBAN, "face_value", value_raw="₹ 1", page=9999, quote="₹ 1 each",
               status="present")  # fmt: skip
    r3 = _fill(ws, URBAN, "registrar", value_raw="Link Intime", page="twelve",
               quote="something else", status="present")  # fmt: skip
    r4 = _fill(ws, URBAN, "promoters", value_raw="Mr X", page=5, status="present")
    r5 = _fill(ws, URBAN, "offer_price", value_raw="₹ 100", page=3, quote="₹ 100", status="done")
    wb.save(xlsx)
    result = import_xlsx(xlsx, tmp_path / "g.jsonl", partial=True)
    by_row = {int(e.split(" ")[1]): e for e in result.errors}
    assert "amount" in by_row[r1]
    assert "₹ 4,720 million" in by_row[r1]  # says how to fix it
    assert "577" in by_row[r2]
    assert "whole number" in by_row[r3]
    assert "quote" in by_row[r3] or "Quote" in by_row[r3]
    assert "Quote is empty" in by_row[r4]
    assert "status" in by_row[r5].lower()
    assert "dropdown" in by_row[r5]
    assert not (tmp_path / "g.jsonl").exists()


def test_partial_import_ignores_untouched_rows(xlsx: Path, tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx)
    _fill(wb["LABELS"], URBAN, "face_value", value_raw="₹ 1", page=9, quote="₹ 1 each",
          status="present")  # fmt: skip
    wb.save(xlsx)
    result = import_xlsx(xlsx, tmp_path / "g.jsonl", partial=True)
    assert (result.written, result.errors, result.unfilled) == (1, [], 109)


def test_complete_import_of_a_fully_filled_sheet(xlsx: Path, tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["LABELS"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7, "not_in_document")
    wb.save(xlsx)
    out = tmp_path / "gold_values.jsonl"
    result = import_xlsx(xlsx, out)
    assert (result.written, result.errors) == (110, [])
    assert validate_file(out, require_complete=True) == []
